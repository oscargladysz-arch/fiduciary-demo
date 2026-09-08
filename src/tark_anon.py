"""
Anonymization tokens: one source of truth
==========================================
The four reference plans are anonymized on every surface. The build refuses
to emit a bundle that leaks a sponsor token, and every test suite checks the
same list. Before this module, the build derived 5 tokens while two test
files listed 4 and never screened EINs or Form 5500 ack ids.

Tokens come from data/plans/*.json identity_private (sponsor, plan_name,
ein, ack_id). When P3-3 moves identities out of the public repo, this module
is the only place that changes.
"""
from __future__ import annotations

from pathlib import Path

from tark_data import load_plan, plan_keys

# generic corporate and plan words that legitimately appear in anonymized
# display labels (for example "tire & rubber manufacturer")
STOP = {"the", "inc", "inc.", "llc", "llp", "co", "co.", "company",
        "corporation", "corp", "corp.", "usa", "us", "group", "and", "of",
        "for", "plan", "trust", "savings", "profit", "sharing",
        "retirement", "employee", "employees", "bargaining", "unit",
        "restaurants", "tire", "rubber", "&", "(psrp)", "401(k)"}


def forbidden_tokens() -> list[str]:
    toks: set[str] = set()
    for k in plan_keys():
        ident = load_plan(k).get("identity_private", {}) or {}
        for field in ("sponsor", "plan_name"):
            for w in str(ident.get(field, "")).split():
                w = w.strip(",.()").lower()
                if w and len(w) > 3 and w not in STOP and not w.startswith("401("):
                    toks.add(w)
        for field in ("ein", "ack_id"):
            v = str(ident.get(field, "")).strip().lower()
            if v:
                toks.add(v)
    return sorted(toks)


def leaks(text: str) -> list[str]:
    """Tokens present in text (lowercased substring match)."""
    low = text.lower()
    return [t for t in forbidden_tokens() if t in low]


def docx_text(path: Path | str) -> str:
    """Plain text of a .docx: paragraphs plus every table cell."""
    prose, provenance = docx_texts(path)
    return prose + ("\n" + provenance if provenance else "")


# a table column whose header is one of these holds provenance set in code
# style (a cited file name, an accession, a URL): the allowlist gate's prose
# rules do not run over it, the forbidden-string list always does (R3-P2-12)
PROVENANCE_HEADERS = ("Source as written", "Accession and EDGAR URL")


def docx_texts(path: Path | str) -> tuple[str, str]:
    """(reader prose, provenance) of a .docx: paragraphs and every table
    cell, with the cells under a provenance header set apart."""
    from docx import Document
    d = Document(str(path))
    prose = [p.text for p in d.paragraphs]
    prov: list[str] = []
    for t in d.tables:
        header = [c.text.strip() for c in t.rows[0].cells] if t.rows else []
        prov_cols = {i for i, h in enumerate(header) if h in PROVENANCE_HEADERS}
        for r, row in enumerate(t.rows):
            for i, c in enumerate(row.cells):
                (prov if (r > 0 and i in prov_cols) else prose).append(c.text)
    return "\n".join(prose), "\n".join(prov)


if __name__ == "__main__":
    print(len(forbidden_tokens()), "forbidden tokens")
